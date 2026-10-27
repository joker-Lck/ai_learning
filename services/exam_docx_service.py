"""
试卷 DOCX 导出 — 可打印标准试卷版式（试题卷 + 答案与解析分页）
"""

import time
from pathlib import Path

from core.logger import error, info

EXPORT_DIR = Path("exports")

TYPE_SECTIONS = [
    ("multiple_choice", "选择题"),
    ("judge", "判断题"),
    ("fill_blank", "填空题"),
    ("short_answer", "简答题"),
]


def _set_run_font(run, name: str = "宋体", size: float = 11, bold: bool = False):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.element.rPr.rFonts.set(qn("w:eastAsia"), name)


class ExamDocxService:
    """试卷导出 DOCX"""

    def export(self, paper: dict, *, with_answers: bool = True) -> dict:
        """
        导出试卷。paper 需含 title/subject/grade/total_score/duration_minutes/questions。
        with_answers=True 时文末附「参考答案与解析」。
        """
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Pt

            questions = paper.get("questions") or []
            if not questions:
                return {"success": False, "message": "试卷没有题目"}

            doc = Document()
            self._setup_page(doc)

            # ── 抬头 ──
            title_p = doc.add_paragraph()
            title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_run_font(title_p.add_run(paper.get("title") or "测试卷"), "黑体", 18, True)

            info_p = doc.add_paragraph()
            info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            meta = (f"学科：{paper.get('subject') or '—'}　　年级：{paper.get('grade') or '—'}　　"
                    f"考试时间：{int(paper.get('duration_minutes') or 60)} 分钟　　"
                    f"满分：{paper.get('total_score') or 100} 分")
            _set_run_font(info_p.add_run(meta), "宋体", 10.5)

            line_p = doc.add_paragraph()
            _set_run_font(line_p.add_run("姓名：____________　班级：____________　得分：____________"),
                          "宋体", 10.5)

            doc.add_paragraph()

            # ── 试题 ──
            seq = 0
            for qtype, label in TYPE_SECTIONS:
                group = [q for q in questions if self._qtype(q) == qtype]
                if not group:
                    continue
                section_p = doc.add_paragraph()
                _set_run_font(section_p.add_run(f"一、{label}" if seq == 0 else
                                                f"{self._cn_num(list(TYPE_SECTIONS).index((qtype, label)) + 1)}、{label}"),
                              "黑体", 13, True)
                total = sum(float(q.get("score") or 0) for q in group)
                hint = f"（共 {len(group)} 小题，每题 {group[0].get('score') or 0} 分，合计 {total:g} 分）"
                _set_run_font(section_p.add_run(hint), "宋体", 9)
                for q in group:
                    seq += 1
                    self._add_question(doc, seq, q, qtype)

            # ── 答案与解析 ──
            if with_answers:
                doc.add_page_break()
                ans_title = doc.add_paragraph()
                ans_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_run_font(ans_title.add_run("参考答案与解析"), "黑体", 16, True)
                for i, q in enumerate(questions, 1):
                    p = doc.add_paragraph()
                    qtype = self._qtype(q)
                    answer = q.get("answer") or "（略）"
                    if qtype == "short_answer":
                        _set_run_font(p.add_run(f"{i}. 【参考答案】"), "黑体", 11, True)
                        ref = doc.add_paragraph()
                        _set_run_font(ref.add_run(q.get("reference_answer") or "（略）"), "宋体", 10.5)
                        points = q.get("scoring_points") or []
                        if points:
                            pp = doc.add_paragraph()
                            _set_run_font(pp.add_run("【评分要点】"), "黑体", 10.5, True)
                            for idx, pt in enumerate(points, 1):
                                bp = doc.add_paragraph()
                                _set_run_font(bp.add_run(f"（{idx}）{pt}"), "宋体", 10.5)
                    else:
                        _set_run_font(p.add_run(f"{i}. 答案：{answer}"), "黑体", 10.5, True)
                    if q.get("explanation"):
                        ep = doc.add_paragraph()
                        _set_run_font(ep.add_run(f"　　解析：{q['explanation']}"), "宋体", 10)

            EXPORT_DIR.mkdir(parents=True, exist_ok=True)
            safe_title = re_sub_filename(paper.get("title") or "试卷")
            filename = f"{safe_title}_{int(time.time())}.docx"
            file_path = EXPORT_DIR / filename
            doc.save(str(file_path))
            info(f"试卷已导出: {file_path}")
            return {"success": True, "message": "导出成功", "filename": filename,
                    "path": str(file_path).replace("\\", "/")}
        except Exception as e:
            error(f"试卷导出失败: {e}")
            return {"success": False, "message": f"导出失败: {e}"}

    # ── 内部 ──

    def _setup_page(self, doc):
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt

        for section in doc.sections:
            section.top_margin = Cm(2.0)
            section.bottom_margin = Cm(2.0)
            section.left_margin = Cm(2.2)
            section.right_margin = Cm(2.2)
        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(11)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    @staticmethod
    def _qtype(q: dict) -> str:
        t = (q.get("type") or q.get("question_type") or "").strip().lower()
        if t in ("multiple_choice", "choice", "single_choice", "选择题"):
            return "multiple_choice"
        if t in ("judge", "true_false", "判断题"):
            return "judge"
        if t in ("fill_blank", "fill", "填空题"):
            return "fill_blank"
        return "short_answer"

    @staticmethod
    def _cn_num(n: int) -> str:
        return "一二三四五六七八九十"[n - 1] if 1 <= n <= 10 else str(n)

    def _add_question(self, doc, seq: int, q: dict, qtype: str):
        from docx.shared import Pt  # noqa: F401

        p = doc.add_paragraph()
        score = q.get("score") or 0
        _set_run_font(p.add_run(f"{seq}. "), "宋体", 11, True)
        _set_run_font(p.add_run(q.get("question") or q.get("question_text") or ""), "宋体", 11)
        _set_run_font(p.add_run(f"（{score:g} 分）"), "宋体", 9)

        if qtype == "multiple_choice":
            for opt in (q.get("options") or []):
                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Pt(18)
                _set_run_font(op.add_run(str(opt)), "宋体", 10.5)
        elif qtype == "short_answer":
            for _ in range(4):
                blank = doc.add_paragraph()
                _set_run_font(blank.add_run("_" * 52), "宋体", 10.5)


def re_sub_filename(name: str) -> str:
    """去掉文件名非法字符"""
    import re

    return re.sub(r'[\\/:*?"<>|\s]+', "_", name).strip("_")[:40] or "试卷"


exam_docx_service = ExamDocxService()
