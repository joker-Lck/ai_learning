"""
教案 DOCX 导出服务
使用 python-docx 生成专业教学设计文档
"""
import os
import uuid
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core.logger import error, info


def _set_cell_shading(cell, color_hex: str):
    """给表格单元格设置背景色"""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)


class DocxLessonService:
    """教案 DOCX 导出服务"""

    def __init__(self):
        self.export_dir = "exports"
        os.makedirs(self.export_dir, exist_ok=True)

    def export(self, lesson_plan: dict) -> dict:
        """
        将教案字典导出为 .docx

        Args:
            lesson_plan: 教案字典（lesson_plan_service.generate_lesson_plan 返回格式）

        Returns:
            {success, file_path, filename, message}
        """
        try:
            doc = Document()
            self._setup_styles(doc)
            self._add_cover(doc, lesson_plan)
            self._add_objectives(doc, lesson_plan)
            self._add_key_difficult(doc, lesson_plan)
            self._add_methods(doc, lesson_plan)
            self._add_process(doc, lesson_plan)
            self._add_board(doc, lesson_plan)
            self._add_homework(doc, lesson_plan)
            self._add_reflection(doc, lesson_plan)

            fname = f"lesson_plan_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.docx"
            fpath = os.path.join(self.export_dir, fname)
            doc.save(fpath)
            info(f"教案 DOCX 导出成功: {fpath}")
            return {"success": True, "file_path": fpath, "filename": fname, "message": "导出成功"}
        except Exception as e:
            error(f"教案 DOCX 导出失败: {e!s}")
            return {"success": False, "message": str(e)}

    # ── 文档结构 ──

    def _setup_styles(self, doc):
        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(12)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.line_spacing = 1.5

    def _add_cover(self, doc, p):
        title = doc.add_heading("", level=0)
        run = title.add_run(p.get("lesson_title", "教学设计"))
        run.font.size = Pt(26)
        run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        run.font.name = "黑体"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        info_items = [
            ("学  科", p.get("subject", "")),
            ("课  题", p.get("topic", "")),
            ("年  级", p.get("grade", "")),
            ("课  时", f"{p.get('duration', 45)} 分钟"),
            ("日  期", datetime.now().strftime("%Y年%m月%d日")),
        ]
        table = doc.add_table(rows=len(info_items), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (label, value) in enumerate(info_items):
            c0 = table.rows[i].cells[0]
            c1 = table.rows[i].cells[1]
            c0.text = label
            c1.text = value
            for cell in (c0, c1):
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(13)
            _set_cell_shading(c0, "EFF6FF")

        doc.add_paragraph()  # 空行

    def _add_section_title(self, doc, text):
        h = doc.add_heading("", level=1)
        run = h.add_run(text)
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
        run.font.name = "黑体"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")

    def _add_objectives(self, doc, p):
        self._add_section_title(doc, "一、教学目标")
        objs = p.get("teaching_objectives", {})
        sections = [("知识与技能", objs.get("knowledge", [])),
                     ("过程与方法", objs.get("ability", [])),
                     ("情感态度与价值观", objs.get("emotion", []))]
        for label, items in sections:
            doc.add_paragraph(f"【{label}】", style="List Bullet")
            for item in items:
                doc.add_paragraph(item, style="List Bullet")

    def _add_key_difficult(self, doc, p):
        self._add_section_title(doc, "二、教学重点与难点")
        doc.add_paragraph("【教学重点】")
        for pt in p.get("key_points", []):
            doc.add_paragraph(pt, style="List Bullet")
        doc.add_paragraph("【教学难点】")
        for pt in p.get("difficult_points", []):
            doc.add_paragraph(pt, style="List Bullet")

    def _add_methods(self, doc, p):
        self._add_section_title(doc, "三、教学方法与教具")
        doc.add_paragraph(f"教学方法：{'、'.join(p.get('teaching_methods', []))}")
        doc.add_paragraph(f"教学教具：{'、'.join(p.get('teaching_aids', []))}")

    def _add_process(self, doc, p):
        self._add_section_title(doc, "四、教学过程")
        process = p.get("teaching_process", [])
        table = doc.add_table(rows=1 + len(process), cols=5)
        table.style = "Table Grid"
        headers = ["教学环节", "时间", "教师活动", "学生活动", "设计意图"]
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            _set_cell_shading(cell, "1E40AF")
            for run in cell.paragraphs[0].runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.bold = True
                run.font.size = Pt(11)

        for r, step in enumerate(process, 1):
            vals = [
                step.get("phase", ""),
                f"{step.get('duration', '')}分钟",
                step.get("teacher_activity", ""),
                step.get("student_activity", ""),
                step.get("design_intent", ""),
            ]
            for c, v in enumerate(vals):
                cell = table.rows[r].cells[c]
                cell.text = str(v)
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(10.5)

    def _add_board(self, doc, p):
        self._add_section_title(doc, "五、板书设计")
        doc.add_paragraph(p.get("board_design", ""))

    def _add_homework(self, doc, p):
        self._add_section_title(doc, "六、课后作业")
        doc.add_paragraph(p.get("homework", ""))

    def _add_reflection(self, doc, p):
        self._add_section_title(doc, "七、教学反思")
        doc.add_paragraph(p.get("teaching_reflection", "（课后填写）"))
        for _ in range(5):
            doc.add_paragraph("_" * 70)


# 全局单例
docx_lesson_service = DocxLessonService()
